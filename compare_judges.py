"""
Compare LLM Judge (internal knowledge) vs LLM Judge V2 (expected answer)

Reads questions from Excel, runs both judges on each question,
outputs a single CSV with both scores for comparison.

Usage:
    python compare_judges.py
    python compare_judges.py --sheet Breast CNS   # specific sheets only
"""

import csv
import json
import os
import re
import time
import argparse
from dataclasses import dataclass
from typing import List, Optional
import boto3
import openpyxl
# from scipy.stats import spearmanr, pearsonr
# from sklearn.metrics import cohen_kappa_score, mean_absolute_error
# import numpy as np

# ── Config ────────────────────────────────────────────────────────────────────
EXCEL_FILE       = "KP_test_questions_REVISED.xlsx"
OUTPUT_FILE      = "judge_comparison.csv"
SKIP_FIRST_SHEET = True
JUDGE_MODEL_ID   = "meta.llama3-70b-instruct-v1:0"
DELAY_BETWEEN_Q  = 2.0

CSV_COLUMNS = [
    "disease_site",
    "topic",
    "question",
    "expected_answer",
    "source_document",
    "rag_answer",
    "retrieved_contexts",
    # V2 scores (with expected answer)
    "v2_accuracy_score",
    "v2_completeness_score",
    "v2_retrieval_adherence_score",
    # Internal scores (without expected answer)
    "internal_accuracy_score",
    "internal_completeness_score",
    "internal_retrieval_adherence_score",
    # Diffs
    "accuracy_diff",
]


# ── Bedrock helper ─────────────────────────────────────────────────────────────
def invoke_llama(client, prompt: str) -> str:
    full_prompt = (
        "<|begin_of_text|><|start_header_id|>system<|end_header_id|>\n\n"
        "You are an expert evaluator. Respond with valid JSON only.<|eot_id|>"
        "<|start_header_id|>user<|end_header_id|>\n\n"
        f"{prompt}<|eot_id|><|start_header_id|>assistant<|end_header_id|>\n\n{{"
    )
    body = json.dumps({
        "prompt": full_prompt,
        "max_gen_len": 512,
        "temperature": 0.1,
        "top_p": 0.9,
    })
    response   = client.invoke_model(modelId=JUDGE_MODEL_ID, body=body)
    generation = json.loads(response["body"].read())["generation"]
    if not generation.strip().startswith("{"):
        generation = "{" + generation
    return generation


def parse_scores(raw: str) -> dict:
    """Parse JSON from LLM, with regex fallback."""
    raw = re.sub(r"```(?:json)?", "", raw).strip()
    if not raw.startswith("{"):
        idx = raw.find("{")
        raw = raw[idx:] if idx != -1 else raw
    if not raw.endswith("}"):
        idx = raw.rfind("}")
        raw = raw[:idx+1] if idx != -1 else raw
    raw = re.sub(r",(\s*[}\]])", r"\1", raw)

    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        scores = {}
        for key in ("accuracy", "completeness", "retrieval_adherence"):
            m = re.search(rf'"{key}".*?"score":\s*(\d)', raw)
            scores[key] = {"score": int(m.group(1)), "reasoning": ""} if m else {"score": 0, "reasoning": "parse_error"}
        return scores


# ── Judge V2 prompt (WITH expected answer) ────────────────────────────────────
def prompt_v2(question, expected_answer, source_document, rag_answer, contexts):
    contexts_text = ""
    for i, ctx in enumerate(contexts[:5], 1):
        contexts_text += f"\n[Context {i}]\n{ctx[:600]}\n"

    return f"""You are evaluating a Radiation Oncology RAG system. Score the generated answer on three criteria (1-5 scale).

## Question:
{question}

## Expected Answer (reference only — not a strict checklist):
{expected_answer}

## Expected Source Document:
{source_document}

## Generated Answer:
{rag_answer}

## Retrieved Contexts:
{contexts_text}

### 1. Accuracy (1-5)
Does the generated answer align with the clinical facts in the expected answer?
- 5: Fully aligned — key clinical facts consistent, no contradictions
- 4: Mostly aligned — minor differences in detail or emphasis
- 3: Partially aligned — captures some key points but misses or contradicts others
- 2: Poorly aligned — significant factual differences from expected answer
- 1: Contradicts expected answer or provides completely different information

### 2. Completeness (1-5)
Does the answer fully address the question asked?
- 5: Comprehensively addresses all aspects
- 4: Addresses most aspects with minor gaps
- 3: Addresses main point but misses notable details
- 2: Incomplete, missing major aspects
- 1: Does not meaningfully address the question

### 3. Retrieval Adherence (1-5)
- 5: No hallucination AND expected source document cited/covered
- 4: No hallucination but source not directly cited
- 3: Mostly grounded but some unsupported claims
- 2: Notable hallucination OR expected source missed
- 1: Largely hallucinated OR completely wrong sources

Return ONLY valid JSON:
{{
    "accuracy":            {{"score": <1-5>, "reasoning": "<one sentence>"}},
    "completeness":        {{"score": <1-5>, "reasoning": "<one sentence>"}},
    "retrieval_adherence": {{"score": <1-5>, "reasoning": "<one sentence>"}}
}}"""


# ── Internal judge prompt (WITHOUT expected answer) ───────────────────────────
def prompt_internal(question, rag_answer, contexts):
    contexts_text = ""
    for i, ctx in enumerate(contexts[:5], 1):
        contexts_text += f"\n[Context {i}]\n{ctx[:600]}\n"

    return f"""You are evaluating a Radiation Oncology RAG system. Score the generated answer on three criteria (1-5 scale).

## Question:
{question}

## Generated Answer:
{rag_answer}

## Retrieved Contexts:
{contexts_text}

### 1. Accuracy (1-5)
How factually correct is the answer based on your medical knowledge?
Be critical and strict in your assessment.

- 5: Every clinical claim is correct and precise with no errors
- 4: Mostly correct but has minor imprecisions or omissions
- 3: Core concept is correct but contains notable factual errors
- 2: Significant factual errors that could affect clinical decisions
- 1: Fundamentally incorrect or contradicts established guidelines
Important: Only give a 5 if you are fully confident every claim
is correct. If you are uncertain about any claim, score 4 or below.

### 2. Completeness (1-5)
Does the answer fully address the question asked?
- 5: Comprehensively addresses all aspects
- 4: Addresses most aspects with minor gaps
- 3: Addresses main point but misses notable details
- 2: Incomplete, missing major aspects
- 1: Does not meaningfully address the question

### 3. Retrieval Adherence (1-5)
Is the answer grounded in the retrieved contexts?
- 5: Fully grounded, no hallucination
- 4: Mostly grounded, minor extrapolation
- 3: Partially grounded with unsupported claims
- 2: Significant unsupported information
- 1: Largely hallucinated

Return ONLY valid JSON:
{{
    "accuracy":            {{"score": <1-5>, "reasoning": "<one sentence>"}},
    "completeness":        {{"score": <1-5>, "reasoning": "<one sentence>"}},
    "retrieval_adherence": {{"score": <1-5>, "reasoning": "<one sentence>"}}
}}"""


# ── Load Excel ─────────────────────────────────────────────────────────────────
def load_questions(ws) -> List[dict]:
    headers = [cell.value for cell in ws[1]]
    questions = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        q = dict(zip(headers, row))
        questions.append({
            "disease_site":    q.get("Disease Site", ""),
            "topic":           q.get("Topic", ""),
            "question":        str(q.get("Question", "") or "").strip(),
            "expected_answer": str(q.get("Expected Answer", "") or ""),
            "source_document": str(q.get("Source Document", "") or ""),
        })
    return [q for q in questions if q["question"]]


def load_completed(csv_path: str) -> set:
    if not os.path.exists(csv_path):
        return set()
    with open(csv_path, "r", encoding="utf-8") as f:
        return {row["question"].strip() for row in csv.DictReader(f)}


# ── Metrics ────────────────────────────────────────────────────────────────────
# def compute_metrics(v2_scores: List[int], internal_scores: List[int], disease_sites: List[str]):
#     print("\n" + "=" * 60)
#     print("COMPARISON METRICS: V2 (ground truth) vs Internal")
#     print("=" * 60)

#     # Overall
#     mae   = mean_absolute_error(v2_scores, internal_scores)
#     corr_s, p_s = spearmanr(v2_scores, internal_scores)
#     corr_p, p_p = pearsonr(v2_scores, internal_scores)

#     # Kappa: discretize to Low/Mid/High
#     def bucket(s):
#         if s <= 2: return 0
#         elif s == 3: return 1
#         else: return 2

#     v2_b  = [bucket(s) for s in v2_scores]
#     int_b = [bucket(s) for s in internal_scores]
#     kappa = cohen_kappa_score(v2_b, int_b)

#     print(f"\nOverall (n={len(v2_scores)}):")
#     print(f"  MAE:               {mae:.3f}")
#     print(f"  Spearman r:        {corr_s:.3f}  (p={p_s:.4f})")
#     print(f"  Pearson r:         {corr_p:.3f}  (p={p_p:.4f})")
#     print(f"  Cohen's Kappa:     {kappa:.3f}")

#     if kappa > 0.6:
#         print("  → Internal judge is RELIABLE for correctness ✅")
#         print("  → Keep correctness weight as proposed (0.3)")
#     elif kappa > 0.4:
#         print("  → Internal judge is MODERATELY reliable ⚠️")
#         print("  → Consider reducing correctness weight to 0.1-0.15")
#     else:
#         print("  → Internal judge is NOT reliable ❌")
#         print("  → Drop correctness from automated score, use RAG proxy instead")

#     # Per-disease breakdown
#     print(f"\nPer-Disease Breakdown:")
#     print(f"  {'Disease':<25} {'n':>4} {'MAE':>6} {'Spearman':>9} {'Kappa':>7}")
#     print(f"  {'-'*55}")

#     unique_diseases = sorted(set(disease_sites))
#     for disease in unique_diseases:
#         idx = [i for i, d in enumerate(disease_sites) if d == disease]
#         if len(idx) < 3:
#             print(f"  {disease:<25} {len(idx):>4}   (insufficient data)")
#             continue
#         d_v2  = [v2_scores[i] for i in idx]
#         d_int = [internal_scores[i] for i in idx]
#         d_mae = mean_absolute_error(d_v2, d_int)
#         d_corr, _ = spearmanr(d_v2, d_int)
#         d_v2_b  = [bucket(s) for s in d_v2]
#         d_int_b = [bucket(s) for s in d_int]
#         try:
#             d_kappa = cohen_kappa_score(d_v2_b, d_int_b)
#         except Exception:
#             d_kappa = float("nan")
#         print(f"  {disease:<25} {len(idx):>4} {d_mae:>6.3f} {d_corr:>9.3f} {d_kappa:>7.3f}")

#     print("=" * 60)
# ── Main ───────────────────────────────────────────────────────────────────────
def run(sheet_names=None):
    from config import default_config
    from agent import AgenticRAGEngine

    wb = openpyxl.load_workbook(EXCEL_FILE)
    sheets = wb.sheetnames[1:] if SKIP_FIRST_SHEET else wb.sheetnames
    if sheet_names:
        sheets = [s for s in sheets if s in sheet_names]

    print(f"Sheets: {sheets}")
    print(f"Output: {OUTPUT_FILE}\n")

    agent  = AgenticRAGEngine(config=default_config)
    session = boto3.Session(profile_name=default_config.aws.profile_name)
    client  = session.client("bedrock-runtime", region_name=default_config.aws.region)

    completed = load_completed(OUTPUT_FILE)
    file_exists = os.path.exists(OUTPUT_FILE)

    # Collect for metrics
    # all_v2_accuracy       = []
    # all_internal_accuracy = []
    # all_disease_sites     = []

    with open(OUTPUT_FILE, "a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=CSV_COLUMNS)
        if not file_exists:
            writer.writeheader()

        for sheet_name in sheets:
            questions = load_questions(wb[sheet_name])
            print(f"\n{'='*60}")
            print(f"Sheet: {sheet_name} ({len(questions)} questions)")
            print(f"{'='*60}")

            for i, q in enumerate(questions, 1):
                question = q["question"]

                if question in completed:
                    print(f"  [{i:2d}] SKIP: {question[:60]}")
                    continue

                print(f"  [{i:2d}/{len(questions)}] {question[:70]}")

                # 1. Get RAG answer
                try:
                    response   = agent.query(question)
                    rag_answer = response.answer
                    contexts   = []
                    if response.task_plan:
                        for task in response.task_plan.sub_tasks:
                            for chunk in task.retrieved_chunks:
                                if chunk.get("text"):
                                    contexts.append(chunk["text"])
                    if not contexts:
                        contexts = [c.chunk_text for c in response.citations if c.chunk_text]
                    contexts = contexts[:8]
                    contexts_json = json.dumps(contexts)  
                except Exception as e:
                    print(f"    RAG failed: {e}")
                    continue

                # 2. Run V2 judge (with expected answer)
                try:
                    raw_v2 = invoke_llama(client, prompt_v2(
                        question, q["expected_answer"],
                        q["source_document"], rag_answer, contexts
                    ))
                    scores_v2 = parse_scores(raw_v2)
                except Exception as e:
                    print(f"    V2 judge failed: {e}")
                    scores_v2 = {"accuracy": {"score": 0}, "completeness": {"score": 0}, "retrieval_adherence": {"score": 0}}

                time.sleep(1)

                # 3. Run internal judge (without expected answer)
                try:
                    raw_int = invoke_llama(client, prompt_internal(
                        question, rag_answer, contexts
                    ))
                    scores_int = parse_scores(raw_int)
                except Exception as e:
                    print(f"    Internal judge failed: {e}")
                    scores_int = {"accuracy": {"score": 0}, "completeness": {"score": 0}, "retrieval_adherence": {"score": 0}}

                v2_acc  = scores_v2.get("accuracy", {}).get("score", 0)
                int_acc = scores_int.get("accuracy", {}).get("score", 0)

                print(f"    V2 accuracy: {v2_acc}/5  |  Internal accuracy: {int_acc}/5  |  diff: {abs(v2_acc - int_acc)}")

                # 4. Write row
                writer.writerow({
                    "disease_site":                    q["disease_site"],
                    "topic":                           q["topic"],
                    "question":                        question,
                    "expected_answer":                 q["expected_answer"],
                    "source_document":                 q["source_document"],
                    "rag_answer":                      rag_answer,
                    "retrieved_contexts":              contexts_json,
                    "v2_accuracy_score":               v2_acc,
                    "v2_completeness_score":           scores_v2.get("completeness", {}).get("score", 0),
                    "v2_retrieval_adherence_score":    scores_v2.get("retrieval_adherence", {}).get("score", 0),
                    "internal_accuracy_score":         int_acc,
                    "internal_completeness_score":     scores_int.get("completeness", {}).get("score", 0),
                    "internal_retrieval_adherence_score": scores_int.get("retrieval_adherence", {}).get("score", 0),
                    "accuracy_diff":                   abs(v2_acc - int_acc),
                })
                f.flush()

                # all_v2_accuracy.append(v2_acc)
                # all_internal_accuracy.append(int_acc)
                # all_disease_sites.append(q["disease_site"])

                time.sleep(DELAY_BETWEEN_Q)

    # # 5. Print metrics
    # if all_v2_accuracy:
    #     compute_metrics(all_v2_accuracy, all_internal_accuracy, all_disease_sites)
    # else:
    #     print("\nNo new results to compute metrics on.")
    #     print("(All questions already completed — check judge_comparison.csv)")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--sheet", nargs="+", default=None)
    args = parser.parse_args()
    run(sheet_names=args.sheet)
