"""
Run once after collecting human evaluation labels.
Must be run before starting evaluation/llm_judge_v3.py.

Steps:
  1. Read Excel (all disease-site sheets)
  2. Average two physician evaluation scores (H1, H2) → ground truth
     Rows missing either physician score are skipped.
  3. Fit Isotonic Regression for correctness_b and answer_quality vs averaged human scores
  4. Add calibrated columns to combined_eval.csv
  5. Output calibration_map.json
  6. Compute Spearman rho (calibrated scores + raw faithfulness/retrieval vs avg human correctness)
     → output weights.json

Outputs
-------
  calibration/calibration_map.json  per-criteria lookup table {score: calibrated_value}
  calibration/weights.json          normalized Spearman weights
  calibration/combined_eval.csv     all rows with raw + calibrated + averaged human scores

Usage
-----
  python calibration/calibration.py --excel Test_question_Human_Eval.xlsx
"""

import argparse
import csv
import json
import os
import sys
import numpy as np
from pathlib import Path

_CALIBRATION_DIR = Path(__file__).parent
import openpyxl
from scipy import stats
from sklearn.isotonic import IsotonicRegression


EXCEL_COLUMNS = {
    "disease_site":     "Disease Site",
    "corr_b_raw":       "correctness b (rag proxy)",
    "quality_raw":      "answer_quality_score(completeness)",
    "faithfulness_raw": "faithfulness_score",
    "retrieval_raw":    "retrieval_adequacy_score",
    "h1_corr":          "Correctness_H1",
    "h1_qual":          "Answer_Quality_H1",
    "h2_corr":          "Correctness_H2",
    "h2_qual":          "Answer_Quality_H2",
}

SKIP_SHEETS = {"Summary + Legend", "Kishan Evaluation Summary"}


def read_excel(excel_path: str) -> list:
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    rows = []
    for sheet_name in wb.sheetnames:
        if sheet_name in SKIP_SHEETS:
            continue
        ws = wb[sheet_name]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            record = dict(zip(headers, row))
            mapped = {"sheet": sheet_name}
            for canonical, excel_col in EXCEL_COLUMNS.items():
                mapped[canonical] = record.get(excel_col)

            # Average two physician scores — None if either is missing
            h1c = mapped.get("h1_corr")
            h2c = mapped.get("h2_corr")
            h1q = mapped.get("h1_qual")
            h2q = mapped.get("h2_qual")

            mapped["human_corr"] = (float(h1c) + float(h2c)) / 2 if (h1c is not None and h2c is not None) else None
            mapped["human_qual"] = (float(h1q) + float(h2q)) / 2 if (h1q is not None and h2q is not None) else None

            rows.append(mapped)
    wb.close()
    print(f"  Read {len(rows)} rows from {len(wb.sheetnames) - len(SKIP_SHEETS)} sheets")
    skipped = sum(1 for r in rows if r["human_corr"] is None or r["human_qual"] is None)
    print(f"  Rows with both physician scores present: {len(rows) - skipped}  (skipped {skipped} missing)")
    return rows


def fit_isotonic(x_vals, y_vals) -> dict:
    """
    Fit isotonic regression and return a lookup table {raw_score: calibrated_value}
    for integer scores 1-5.
    Isotonic guarantees monotone mapping: higher raw score → higher calibrated score.
    y_vals are averaged physician scores and may be non-integer (e.g. 3.5).
    """
    X = np.array(x_vals)
    Y = np.array(y_vals)
    ir = IsotonicRegression(out_of_bounds="clip").fit(X, Y)
    lookup = {}
    for score in range(1, 6):
        lookup[str(score)] = round(float(ir.predict([score])[0]), 4)
    return lookup


def apply_isotonic(raw_score, lookup: dict) -> float:
    key = str(int(raw_score))
    return float(lookup.get(key, raw_score))


def compute_weights(rows, calibration_map):
    valid = [
        r for r in rows
        if all(r[k] is not None for k in
               ["human_corr", "corr_b_raw", "quality_raw",
                "faithfulness_raw", "retrieval_raw"])
    ]
    print(f"    Using {len(valid)} rows for weight computation")
    human = np.array([float(r["human_corr"]) for r in valid])

    criteria_vals = {
        "retrieval_adequacy": np.array([float(r["retrieval_raw"])    for r in valid]),
        "faithfulness":       np.array([float(r["faithfulness_raw"]) for r in valid]),
        "correctness":        np.array([
            apply_isotonic(float(r["corr_b_raw"]), calibration_map["correctness"])
            for r in valid
        ]),
        "answer_quality":     np.array([
            apply_isotonic(float(r["quality_raw"]), calibration_map["answer_quality"])
            for r in valid
        ]),
    }

    rhos = {}
    for name, vals in criteria_vals.items():
        rho, _ = stats.spearmanr(vals, human)
        rhos[name] = abs(rho)
        print(f"    rho({name:<20} vs avg physician correctness) = {rho:.4f}")

    total = sum(rhos.values())
    weights = {k: round(v / total, 4) for k, v in rhos.items()}
    # Ensure exact sum = 1.0
    diff = round(1.0 - sum(weights.values()), 4)
    weights[max(weights, key=weights.get)] = round(
        weights[max(weights, key=weights.get)] + diff, 4
    )
    return weights


def write_json(path, data):
    with open(path, "w") as f:
        json.dump(data, f, indent=2)
    print(f"    Wrote {path}")


def write_combined_csv(output_dir, rows, calibration_map):
    path = os.path.join(output_dir, "combined_eval.csv")
    fieldnames = [
        "sheet", "disease_site",
        "corr_b_raw", "corr_b_calibrated",
        "quality_raw", "quality_calibrated",
        "faithfulness_raw", "retrieval_raw",
        "h1_corr", "h1_qual",
        "h2_corr", "h2_qual",
        "human_corr_avg", "human_qual_avg",
    ]
    with open(path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for r in rows:
            row_out = {
                "sheet":            r["sheet"],
                "disease_site":     r["disease_site"],
                "corr_b_raw":       r["corr_b_raw"],
                "quality_raw":      r["quality_raw"],
                "faithfulness_raw": r["faithfulness_raw"],
                "retrieval_raw":    r["retrieval_raw"],
                "h1_corr":          r["h1_corr"],
                "h1_qual":          r["h1_qual"],
                "h2_corr":          r["h2_corr"],
                "h2_qual":          r["h2_qual"],
                "human_corr_avg":   r["human_corr"],
                "human_qual_avg":   r["human_qual"],
            }
            for crit, raw_key, cal_key in [
                ("correctness",    "corr_b_raw",  "corr_b_calibrated"),
                ("answer_quality", "quality_raw", "quality_calibrated"),
            ]:
                raw = r.get(raw_key)
                try:
                    row_out[cal_key] = apply_isotonic(
                        float(raw), calibration_map[crit]
                    )
                except (TypeError, ValueError):
                    row_out[cal_key] = ""
            writer.writerow(row_out)
    print(f"    Wrote {path}  ({len(rows)} rows)")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--excel",      required=True)
    parser.add_argument("--output_dir", default=_CALIBRATION_DIR)
    args = parser.parse_args()

    os.makedirs(args.output_dir, exist_ok=True)

    print("\n── Step 1: Read Excel ────────────────────────────────")
    rows = read_excel(args.excel)

    print("\n── Step 2: Fit Isotonic Regression ───────────────────")
    calibration_map = {}
    for crit, raw_key, human_key in [
        ("correctness",    "corr_b_raw",  "human_corr"),
        ("answer_quality", "quality_raw", "human_qual"),
    ]:
        paired = [
            (float(r[raw_key]), float(r[human_key]))
            for r in rows
            if r[raw_key] is not None and r[human_key] is not None
        ]
        if len(paired) < 5:
            print(f"  ERROR: {crit} only {len(paired)} labeled rows — need at least 5")
            sys.exit(1)
        xs, ys = zip(*paired)
        lookup = fit_isotonic(list(xs), list(ys))
        calibration_map[crit] = lookup
        print(f"  {crit:<20} n={len(paired)}  lookup={lookup}")

    print("\n── Step 3: Spearman weights ──────────────────────────")
    weights = compute_weights(rows, calibration_map)
    print(f"\n  Final weights:")
    for k, v in weights.items():
        print(f"    {k:<20} {v:.4f}")

    print("\n── Step 4: Write outputs ─────────────────────────────")
    write_json(os.path.join(args.output_dir, "calibration_map.json"), calibration_map)
    write_json(os.path.join(args.output_dir, "weights.json"), weights)
    write_combined_csv(args.output_dir, rows, calibration_map)

    print("── Done ──────────────────────────────────────────────")


if __name__ == "__main__":
    main()
