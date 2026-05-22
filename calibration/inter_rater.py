"""
Inter-rater analysis between two physician evaluators (H1 and H2).

For each criterion (Correctness, Answer Quality):
  - Spearman rho: do the two raters rank answers the same way?
  - Wilcoxon signed-rank test: does one rater systematically score higher?

Rows with any missing H1 or H2 value are excluded.

Usage
-----
  python calibration/inter_rater.py --excel Test_question_Human_Eval.xlsx
"""

import argparse
import csv
from pathlib import Path

import numpy as np
import openpyxl
from scipy import stats

SKIP_SHEETS = {"Summary + Legend", "Kishan Evaluation Summary"}

CRITERIA = [
    # (label, H1 col name, H2 col name)
    ("Correctness",    "Correctness_H1",    "Correctness_H2"),
    ("Answer_Quality", "Answer_Quality_H1", "Answer_Quality_H2"),
]


def read_paired_scores(excel_path: str) -> dict:
    """
    Returns {criterion_label: {"h1": [...], "h2": [...], "sheet": [...]}}
    Only includes rows where both H1 and H2 are non-null for that criterion.
    """
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    data = {label: {"h1": [], "h2": [], "sheet": []} for label, _, _ in CRITERIA}

    for sheet_name in wb.sheetnames:
        if sheet_name in SKIP_SHEETS:
            continue
        ws = wb[sheet_name]
        all_rows = list(ws.iter_rows(min_row=1, values_only=True))
        if not all_rows:
            continue
        hdrs = all_rows[0]
        col = {h: i for i, h in enumerate(hdrs) if h is not None}

        for label, h1_col, h2_col in CRITERIA:
            if h1_col not in col or h2_col not in col:
                continue
            for row in all_rows[1:]:
                if not any(row):
                    continue
                v1 = row[col[h1_col]]
                v2 = row[col[h2_col]]
                if v1 is None or v2 is None:
                    continue
                try:
                    data[label]["h1"].append(float(v1))
                    data[label]["h2"].append(float(v2))
                    data[label]["sheet"].append(sheet_name)
                except (ValueError, TypeError):
                    continue

    wb.close()
    return data


def analyze(data: dict) -> list[dict]:
    """
    For each criterion, compute:
      - n (pairs used)
      - H1 mean, median
      - H2 mean, median
      - mean diff (H1 - H2)
      - Spearman rho + p-value
      - Wilcoxon statistic + p-value
    """
    results = []
    for label, vals in data.items():
        h1 = np.array(vals["h1"])
        h2 = np.array(vals["h2"])
        n = len(h1)

        if n < 2:
            print(f"  WARNING: {label} has only {n} paired rows, skipping.")
            continue

        # Spearman rho: rank correlation — do they agree on ordering?
        rho, rho_p = stats.spearmanr(h1, h2)

        # Wilcoxon signed-rank test: is there a systematic shift in scores?
        # Tests H0: median(H1 - H2) = 0
        diffs = h1 - h2
        if np.all(diffs == 0):
            wilcox_stat, wilcox_p = np.nan, np.nan
        else:
            wilcox_stat, wilcox_p = stats.wilcoxon(h1, h2)

        results.append({
            "criterion":        label,
            "n_pairs":          n,
            "H1_mean":          round(float(h1.mean()), 4),
            "H1_median":        float(np.median(h1)),
            "H2_mean":          round(float(h2.mean()), 4),
            "H2_median":        float(np.median(h2)),
            "mean_diff_H1_H2":  round(float(diffs.mean()), 4),
            "spearman_rho":     round(float(rho), 4),
            "spearman_p":       f"{float(rho_p):.4e}",
            "wilcoxon_stat":    round(float(wilcox_stat), 2) if not np.isnan(wilcox_stat) else "",
            "wilcoxon_p":       f"{float(wilcox_p):.4e}" if not np.isnan(wilcox_p) else "",
            "systematic_bias":  "yes" if (not np.isnan(wilcox_p) and wilcox_p < 0.05) else "no",
        })

    return results


def print_results(results: list[dict]):
    for r in results:
        print(f"\n── {r['criterion']} (n={r['n_pairs']}) ─────────────────────────")
        print(f"  H1: mean={r['H1_mean']:.3f}  median={r['H1_median']}")
        print(f"  H2: mean={r['H2_mean']:.3f}  median={r['H2_median']}")
        print(f"  mean diff (H1 - H2) = {r['mean_diff_H1_H2']:+.3f}")
        print(f"  Spearman rho = {r['spearman_rho']:.4f}  (p={r['spearman_p']})")
        print(f"  Wilcoxon     stat={r['wilcoxon_stat']}  p={r['wilcoxon_p']}")
        print(f"  Systematic bias: {r['systematic_bias']}")


def write_csv(results: list[dict], output_path: str):
    fieldnames = [
        "criterion", "n_pairs",
        "H1_mean", "H1_median", "H2_mean", "H2_median",
        "mean_diff_H1_H2",
        "spearman_rho", "spearman_p",
        "wilcoxon_stat", "wilcoxon_p",
        "systematic_bias",
    ]
    with open(output_path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(results)
    print(f"\n  Wrote {output_path}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--excel",  required=True, help="Path to Test_question_Human_Eval.xlsx")
    parser.add_argument("--output", default="calibration/inter_rater_results.csv")
    args = parser.parse_args()

    print(f"\n── Reading {args.excel} ──────────────────────────────")
    data = read_paired_scores(args.excel)

    for label, vals in data.items():
        print(f"  {label}: {len(vals['h1'])} paired rows (missing excluded)")

    print("\n── Results ───────────────────────────────────────────")
    results = analyze(data)
    print_results(results)

    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    write_csv(results, args.output)


if __name__ == "__main__":
    main()
