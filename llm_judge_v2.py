"""
Automated RAG Evaluation Pipeline

Logic:
  For each disease site tab in the Excel file:
    For each question (skipping already-completed ones):
      1. Call agent.query(question) → RAG answer + retrieved chunks
      2. Call LLM judge with question + expected_answer + retrieved_chunks → 3 scores
      3. Append result to per-disease-site CSV immediately
"""

import csv
import json
import os
import re
import time
import argparse
from dataclasses import dataclass
from typing import List, Optional
import boto3
import openpyxl
from config import default_config, create_boto3_session
from agent import AgenticRAGEngine


EXCEL_FILE      = "KP_test_questions_REVISED.xlsx"
OUTPUT_DIR      = "eval_results"          # one CSV per disease site goes here
SKIP_FIRST_SHEET = True                   # skip "Summary + Legend" tab
JUDGE_MODEL_ID  = "meta.llama3-70b-instruct-v1:0"
DELAY_BETWEEN_Q = 2.0                     # seconds between questions

CSV_COLUMNS = [
    "disease_site",
    "topic",
    "question",
    "expected_answer",
    "source_document",
    "rag_answer",
    "accuracy_score",
    "completeness_score",
    "retrieval_adherence_score",
]

@dataclass
class EvalResult:
    disease_site: str
    topic: str
    question: str
    expected_answer: str
    source_document: str
    rag_answer: str
    accuracy_score: int
    completeness_score: int
    retrieval_adherence_score: int


class LLMJudgeV2:
    def __init__(self):
        self.session = create_boto3_session()
        self.client  = self.session.client(
            "bedrock-runtime",
            region_name=default_config.aws.region
        )

    def _build_prompt(
        self,
        question: str,
        expected_answer: str,
        source_document: str,
        rag_answer: str,
        retrieved_contexts: List[str],
    ) -> str:
        contexts_text = ""
        for i, ctx in enumerate(retrieved_contexts[:5], 1):
            contexts_text += f"\n[Context {i}]\n{ctx[:600]}\n"

        return f"""You are evaluating a Radiation Oncology RAG system. Score the generated answer on three criteria (1-5 scale).

## Question:
{question}

## Expected Answer (reference only — not a strict checklist):
{expected_answer}

## Expected Source Document:
{source_document}

## Generated Answer:
{rag_answer}

## Retrieved Contexts:
{contexts_text}

---

## Scoring Criteria:

### 1. Accuracy (1-5)
Use the expected answer as a reference point, not a strict checklist.
Does the generated answer align with the clinical facts in the expected answer?
- 5: Fully aligned — key clinical facts consistent, no contradictions
- 4: Mostly aligned — minor differences in detail or emphasis
- 3: Partially aligned — captures some key points but misses or contradicts others
- 2: Poorly aligned — significant factual differences from expected answer
- 1: Contradicts expected answer or provides completely different information

### 2. Completeness (1-5)
Does the answer fully address the question asked?
- 5: Comprehensively addresses all aspects of the question
- 4: Addresses most aspects with minor gaps
- 3: Addresses main point but misses notable details
- 2: Incomplete, missing major aspects of the question
- 1: Does not meaningfully address the question

### 3. Retrieval Adherence (1-5)
Two components: (a) no hallucination, (b) expected source document covered.
- 5: No hallucination AND expected source document cited or content clearly covered
- 4: No hallucination but expected source not directly cited (content partially covered)
- 3: Mostly grounded but some unsupported claims, or source only partially covered
- 2: Notable hallucination OR expected source document completely missed
- 1: Largely hallucinated OR completely wrong sources used

---

Return ONLY valid JSON, no markdown, no extra text:
{{
    "accuracy":            {{"score": <1-5>, "reasoning": "<one sentence>"}},
    "completeness":        {{"score": <1-5>, "reasoning": "<one sentence>"}},
    "retrieval_adherence": {{"score": <1-5>, "reasoning": "<one sentence>"}}
}}"""

    def _invoke(self, prompt: str) -> str:
        """Call Llama3 on Bedrock."""
        full_prompt = (
            "<|begin_of_text|><|start_header_id|>system<|end_header_id|>\n\n"
            "You are an expert evaluator. Respond with valid JSON only.<|eot_id|>"
            "<|start_header_id|>user<|end_header_id|>\n\n"
            f"{prompt}<|eot_id|><|start_header_id|>assistant<|end_header_id|>\n\n{{"
        )
        body = json.dumps({
            "prompt":       full_prompt,
            "max_gen_len":  512,
            "temperature":  0.1,
            "top_p":        0.9,
        })
        response      = self.client.invoke_model(modelId=JUDGE_MODEL_ID, body=body)
        generation    = json.loads(response["body"].read())["generation"]
        if not generation.strip().startswith("{"):
            generation = "{" + generation
        return generation

    def _parse(self, raw: str) -> dict:
        """Parse JSON from LLM response, with fallback regex."""
        # Clean markdown fences
        raw = re.sub(r"```(?:json)?", "", raw).strip()
        if not raw.startswith("{"):
            idx = raw.find("{")
            raw = raw[idx:] if idx != -1 else raw
        if not raw.endswith("}"):
            idx = raw.rfind("}")
            raw = raw[:idx+1] if idx != -1 else raw
        raw = re.sub(r",(\s*[}\]])", r"\1", raw)  # trailing commas

        try:
            return json.loads(raw)
        except json.JSONDecodeError:
            # Regex fallback
            scores = {}
            for key in ("accuracy", "completeness", "retrieval_adherence"):
                m = re.search(rf'"{key}".*?"score":\s*(\d)', raw)
                scores[key] = {"score": int(m.group(1)), "reasoning": ""} if m else {"score": 0, "reasoning": ""}
            return scores

    def score(
        self,
        question: str,
        expected_answer: str,
        source_document: str,
        rag_answer: str,
        retrieved_contexts: List[str],
    ) -> dict:
        """Returns {accuracy_score, completeness_score, retrieval_adherence_score}."""
        prompt = self._build_prompt(
            question, expected_answer, source_document, rag_answer, retrieved_contexts
        )
        try:
            raw    = self._invoke(prompt)
            parsed = self._parse(raw)
            return {
                "accuracy_score":                parsed.get("accuracy",            {}).get("score", 0),
                "completeness_score":            parsed.get("completeness",        {}).get("score", 0),
                "retrieval_adherence_score":     parsed.get("retrieval_adherence", {}).get("score", 0),
                "accuracy_reasoning":            parsed.get("accuracy",            {}).get("reasoning", ""),
                "completeness_reasoning":        parsed.get("completeness",        {}).get("reasoning", ""),
                "retrieval_adherence_reasoning": parsed.get("retrieval_adherence", {}).get("reasoning", ""),
            }
        except Exception as e:
            print(f"\n    Judge failed: {e}")
            return {
                "accuracy_score": 0, "completeness_score": 0, "retrieval_adherence_score": 0,
                "accuracy_reasoning": "", "completeness_reasoning": "", "retrieval_adherence_reasoning": "",
            }


# Helpers
def load_questions_from_sheet(ws) -> List[dict]:
    """Read all question rows from an Excel worksheet."""
    headers = [cell.value for cell in ws[1]]
    questions = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):  # skip empty rows
            continue
        q = dict(zip(headers, row))
        questions.append({
            "disease_site":    q.get("Disease Site", ""),
            "topic":           q.get("Topic", ""),
            "question":        q.get("Question", "") or "",
            "expected_answer": q.get("Expected Answer", "") or "",
            "source_document": q.get("Source Document", "") or "",
        })
    return [q for q in questions if q["question"].strip()]


def load_completed_questions(csv_path: str) -> set:
    """Return set of question strings already in the CSV (for resume)."""
    if not os.path.exists(csv_path):
        return set()
    with open(csv_path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return {row["question"].strip() for row in reader}


def extract_retrieved_contexts(response) -> List[str]:
    """Pull retrieved chunk texts out of AgentResponse."""
    contexts = []
    meta = response.metadata or {}
    # agent stores them in task_plan sub_tasks
    if response.task_plan:
        for task in response.task_plan.sub_tasks:
            for chunk in task.retrieved_chunks:
                text = chunk.get("text", "")
                if text:
                    contexts.append(text)
    if not contexts:
        # fallback: citations
        for c in response.citations:
            if c.chunk_text:
                contexts.append(c.chunk_text)
    return contexts[:8]  # cap at 8 to keep prompt manageable


# Main evaluation loop
def run_evaluation(sheet_names: Optional[List[str]] = None, output_dir: str = OUTPUT_DIR):
    os.makedirs(output_dir, exist_ok=True)

    # Load Excel
    wb = openpyxl.load_workbook(EXCEL_FILE)
    all_sheets = wb.sheetnames
    if SKIP_FIRST_SHEET:
        all_sheets = all_sheets[1:]
    if sheet_names:
        all_sheets = [s for s in all_sheets if s in sheet_names]

    print(f"\nSheets to evaluate: {all_sheets}")
    print(f"Output directory:   {output_dir}/\n")

    # Init agent and judge (once, shared across all sheets)
    print("Initializing RAG agent...")
    agent = AgenticRAGEngine(config=default_config)
    judge = LLMJudgeV2()
    print("Agent ready.\n")

    total_done  = 0
    total_skip  = 0
    total_error = 0

    for sheet_name in all_sheets:
        ws         = wb[sheet_name]
        questions  = load_questions_from_sheet(ws)
        csv_path   = os.path.join(output_dir, f"{sheet_name}.csv")
        completed  = load_completed_questions(csv_path)

        print(f"{'='*60}")
        print(f"Sheet: {sheet_name}  ({len(questions)} questions, {len(completed)} already done)")
        print(f"{'='*60}")

        # Open CSV in append mode — safe for resume
        file_exists = os.path.exists(csv_path)
        with open(csv_path, "a", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=CSV_COLUMNS)
            if not file_exists:
                writer.writeheader()  # only write header for new files

            for i, q in enumerate(questions, 1):
                question = q["question"].strip()

                # Skip already completed
                if question in completed:
                    print(f"  [{i:2d}/{len(questions)}] SKIP (done): {question[:70]}")
                    total_skip += 1
                    continue
                print(f"  [{i:2d}/{len(questions)}] {question}")

                # 1. RAG answer
                try:
                    response    = agent.query(question)
                    rag_answer  = response.answer
                    contexts    = extract_retrieved_contexts(response)
                    print(f"\nAnswer: {rag_answer}")
                except Exception as e:
                    print(f"\nRAG FAILED: {e}")
                    total_error += 1
                    continue

                # 2.LLM judge scores
                scores = judge.score(
                    question        = question,
                    expected_answer = q["expected_answer"],
                    source_document = q["source_document"],
                    rag_answer      = rag_answer,
                    retrieved_contexts = contexts,
                )
                print(f"\nScores:")
                print(f" - Accuracy ({scores['accuracy_score']}/5): {scores['accuracy_reasoning']}")
                print(f" - Completeness ({scores['completeness_score']}/5): {scores['completeness_reasoning']}")
                print(f" - Retrieval ({scores['retrieval_adherence_score']}/5): {scores['retrieval_adherence_reasoning']}")

                # 3. Write to CSV immediately
                writer.writerow({
                    "disease_site":             q["disease_site"],
                    "topic":                    q["topic"],
                    "question":                 question,
                    "expected_answer":          q["expected_answer"],
                    "source_document":          q["source_document"],
                    "rag_answer":               rag_answer,
                    "accuracy_score":           scores["accuracy_score"],
                    "completeness_score":       scores["completeness_score"],
                    "retrieval_adherence_score":scores["retrieval_adherence_score"],
                })
                f.flush()  # write to disk immediately
                total_done += 1

                time.sleep(DELAY_BETWEEN_Q)

        print(f"  → Saved to {csv_path}\n")

    print(f"\n{'='*60}")
    print(f"Evaluation complete.")
    print(f"  Done:    {total_done}")
    print(f"  Skipped: {total_skip} (already completed)")
    print(f"  Errors:  {total_error}")
    print(f"  CSVs in: {os.path.abspath(output_dir)}/")
    print(f"{'='*60}\n")



if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Automated RAG evaluation pipeline")
    parser.add_argument(
        "--sheet", nargs="+", default=None,
        help="Run specific sheet(s) only, e.g. --sheet Breast CNS"
    )
    args = parser.parse_args()

    run_evaluation(sheet_names=args.sheet)