"""Validate the scenario workbook and export development-only JSONL files.

This module has no AWS dependencies and makes no model calls.
"""

import argparse
from collections import Counter
from contextlib import closing
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import random
import sys

import openpyxl


ROOT = Path(__file__).resolve().parent
DEFAULT_WORKBOOK = ROOT / "data" / "RadOnc-Scenario-Bank-v2-Combined-Excel.xlsx"
VIGNETTE = "Vignette (model input)"
QUESTION = "Clinical Question (model input)"
PLAN = "Attending Plan (reference – never shown to model)"
CORE_FIELDS = ("Case ID", "Patient ID", VIGNETTE, QUESTION, PLAN)
REQUIRED_COLUMNS = CORE_FIELDS + (
    "Split", "Worked-Example Eligible", "LLM Use", "Disease Site", "Subsite",
    "Setting", "Prior RT", "Technique", "Encounter Type",
)
REQUIRED_VALUES = CORE_FIELDS + (
    "Split", "Worked-Example Eligible", "LLM Use", "Disease Site", "Setting",
    "Prior RT", "Encounter Type",
)
CLINICAL_METADATA_FIELDS = {
    "disease_site": "Disease Site",
    "subsite": "Subsite",
    "setting": "Setting",
    "prior_rt": "Prior RT",
    "technique": "Technique",
    "encounter_type": "Encounter Type",
}
PATIENT_LINKAGE_FIELDS = {
    "patient_id": "Patient ID",
    "encounter_number": "Encounter #",
    "first_encounter_at_site": "First Encounter (site)",
}
GOVERNANCE_FIELDS = {
    "split": "Split",
    "worked_example_eligible": "Worked-Example Eligible",
    "llm_use": "LLM Use",
}
PROVENANCE_FIELDS = {
    "batch": "Batch",
    "source_code": "Source Code",
}


def blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def read_records(path):
    """Read only Scenario Bank, preserving column names and supplied values."""
    with closing(openpyxl.load_workbook(path, read_only=True, data_only=False)) as workbook:
        if "Scenario Bank" not in workbook.sheetnames:
            raise ValueError("Workbook is missing the Scenario Bank tab.")
        rows = workbook["Scenario Bank"].iter_rows(values_only=True)
        headers = list(next(rows, ()))
        if not headers or any(blank(h) for h in headers):
            raise ValueError("Scenario Bank must have a non-empty header for each column.")
        if len(set(headers)) != len(headers):
            raise ValueError("Scenario Bank contains duplicate column headers.")
        missing = sorted(set(REQUIRED_COLUMNS) - set(headers))
        if missing:
            raise ValueError("Missing required columns: " + ", ".join(missing))
        records = []
        for row_number, values in enumerate(rows, start=2):
            if all(blank(v) for v in values):
                continue
            if any(isinstance(v, str) and v.startswith("=") for v in values):
                raise ValueError(f"Row {row_number}: formulas are not supported in scenario records.")
            records.append(dict(zip(headers, values)))
    return records


def validate_records(records):
    """Reject invalid labels, missing essentials, duplicates and patient leakage."""
    errors = []
    seen = set()
    patients = {"Development": set(), "Locked test": set()}
    expected = {
        "Development": ("Yes", "Worked-example index"),
        "Locked test": ("No", "Evaluation only – never index"),
    }
    if not records:
        errors.append("Scenario Bank contains no records.")
    for row_number, record in enumerate(records, start=2):
        prefix = f"Record {row_number - 1}"
        for field in REQUIRED_VALUES:
            if blank(record.get(field)):
                errors.append(f"{prefix}: missing {field}.")
        case_id = record.get("Case ID")
        if not blank(case_id):
            if case_id in seen:
                errors.append(f"{prefix}: duplicate Case ID.")
            seen.add(case_id)
        split = record.get("Split")
        if split not in expected:
            errors.append(f"{prefix}: invalid Split label.")
        else:
            actual = (record.get("Worked-Example Eligible"), record.get("LLM Use"))
            if actual != expected[split]:
                errors.append(f"{prefix}: eligibility or LLM Use conflicts with Split.")
            if not blank(record.get("Patient ID")):
                patients[split].add(record["Patient ID"])
    if patients["Development"] & patients["Locked test"]:
        errors.append("Patient overlap detected between Development and Locked test.")
    if errors:
        raise ValueError("Validation failed:\n" + "\n".join(errors))


def development_examples(records):
    """Export the approved field groups without mixing metadata roles."""
    return [
        {
            "case_id": r["Case ID"],
            "vignette": r[VIGNETTE],
            "clinical_question": r[QUESTION],
            "search_text": f"{r[VIGNETTE]}\n\n{r[QUESTION]}",
            "attending_plan": r[PLAN],
            "label": "Institutional example – not evidence",
            "metadata": {
                key: r[column]
                for key, column in CLINICAL_METADATA_FIELDS.items()
            },
            "patient_linkage": {
                key: r.get(column)
                for key, column in PATIENT_LINKAGE_FIELDS.items()
            },
            "governance": {
                key: r[column]
                for key, column in GOVERNANCE_FIELDS.items()
            },
            "provenance": {
                key: r.get(column)
                for key, column in PROVENANCE_FIELDS.items()
            },
        }
        for r in sorted(records, key=lambda r: r["Case ID"])
        if r["Split"] == "Development" and r["Worked-Example Eligible"] == "Yes"
    ]


def baseline_inputs(examples, size, seed):
    """Select reproducibly; model inputs contain no plans or routing metadata."""
    if size < 1 or size > len(examples):
        raise ValueError(f"Baseline size must be between 1 and {len(examples)}.")
    selected = random.Random(seed).sample(sorted(examples, key=lambda r: r["case_id"]), size)
    return [
        {key: r[key] for key in ("case_id", "vignette", "clinical_question")}
        for r in sorted(selected, key=lambda r: r["case_id"])
    ]


def write_jsonl(path, records):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False) + "\n")


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--baseline-size", type=int, default=10)
    parser.add_argument("--seed", type=int, default=42)
    args = parser.parse_args(argv)
    try:
        records = read_records(args.workbook)
        validate_records(records)
        examples = development_examples(records)
        baseline = baseline_inputs(examples, args.baseline_size, args.seed)
        summary = {
            "validation": "passed",
            "generated_at_utc": datetime.now(timezone.utc).isoformat(),
            "source_filename": args.workbook.name,
            "source_sha256": hashlib.sha256(args.workbook.read_bytes()).hexdigest(),
            "total_scenarios": len(records),
            "unique_patients": len({r["Patient ID"] for r in records}),
            "split_counts": dict(Counter(r["Split"] for r in records)),
            "disease_site_counts": dict(sorted(Counter(r["Disease Site"] for r in records).items())),
            "patient_overlap_count": 0,
            "development_examples_exported": len(examples),
            "locked_test_records_exported": 0,
            "baseline_size": len(baseline),
            "baseline_seed": args.seed,
        }
        # Validate and construct all results before writing any export.
        processed = ROOT / "data" / "processed"
        write_jsonl(processed / "development_examples.jsonl", examples)
        write_jsonl(processed / "baseline_inputs.jsonl", baseline)
        output = ROOT / "outputs" / "data_summary.json"
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(summary, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    except (OSError, ValueError, KeyError) as exc:
        print(f"Loader error: {exc}\nNo successful run reported. Existing exports may be stale.", file=sys.stderr)
        return 1
    print(json.dumps(summary, indent=2, ensure_ascii=False))
    print(f"Development exports: {processed}")
    print(f"Summary: {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
